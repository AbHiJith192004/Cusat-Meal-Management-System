#!/usr/bin/env python3
"""Set campus_location from the registration sheet's Lakeside column.

    ops/set_lakeside.py "CUSAT App reg details.xlsx"            # dry run
    ops/set_lakeside.py "CUSAT App reg details.xlsx" --apply    # write it

Reads the sheet, works out who is Lakeside, compares that against what the
server currently holds, and prints the difference. Nothing is written without
--apply, and even then only the students whose campus actually differs are
touched -- the endpoint rejects a no-op change anyway.

The sheet contains resubmissions: several students filled the form more than
once. The LATEST row per student id wins, exactly as the import did, so an
early "No" cannot override a later "Yes".

Credentials are prompted for and go straight to the server. Campus decides the
billing rate, so every change is audited server-side with the reason below.
"""
from __future__ import annotations

import getpass
import json
import sys
import urllib.error
import urllib.request

BASE = "https://messconnect-uyeus.ondigitalocean.app/api/v1"
REASON = "Campus corrected from the CUSAT registration sheet's Lakeside column"


def call(method: str, path: str, token: str | None = None, body: dict | None = None):
    data = json.dumps(body).encode() if body is not None else None
    req = urllib.request.Request(BASE + path, data=data, method=method)
    req.add_header("Content-Type", "application/json")
    if token:
        req.add_header("Authorization", "Bearer " + token)
    try:
        with urllib.request.urlopen(req, timeout=30) as r:
            return json.loads(r.read())
    except urllib.error.HTTPError as e:
        detail = e.read().decode(errors="replace")
        raise SystemExit(f"{method} {path} failed: HTTP {e.code}\n{detail}")


def lakeside_ids_from_sheet(path: str) -> tuple[dict[str, str], int]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise SystemExit("openpyxl is needed: pip install openpyxl")

    ws = load_workbook(path, read_only=True, data_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(c).strip().lower() if c is not None else "" for c in rows[0]]
    try:
        c_id = header.index("student id")
        c_name = header.index("full name")
        c_lake = header.index("lakeside")
    except ValueError:
        raise SystemExit(f"Unexpected columns. Found: {header}")

    latest: dict[str, tuple] = {}
    for row in rows[1:]:
        sid = str(row[c_id]).strip() if row[c_id] is not None else ""
        if not sid:
            continue
        stamp = row[0]
        previous = latest.get(sid)
        if previous is None or (stamp is not None and previous[0] is not None and stamp >= previous[0]):
            latest[sid] = (stamp, row)

    wanted = {
        sid: str(row[c_name] or "").strip()
        for sid, (_, row) in latest.items()
        if str(row[c_lake]).strip().lower() == "yes"
    }
    return wanted, len(latest)


def main() -> int:
    if len(sys.argv) < 2:
        raise SystemExit(__doc__)
    sheet = sys.argv[1]
    apply = "--apply" in sys.argv

    wanted, unique = lakeside_ids_from_sheet(sheet)
    print(f"sheet: {unique} unique student ids, {len(wanted)} marked Lakeside\n")

    reg = input("Super admin registration number: ").strip()
    password = getpass.getpass("Password: ")
    token = call("POST", "/auth/login", body={
        "registration_number": reg, "password": password})["data"]["access_token"]

    roll: list[dict] = []
    for page in range(1, 100):
        rows = call("GET", f"/admin/students?page={page}&per_page=100", token)["data"]
        roll.extend(rows)
        if len(rows) < 100:
            break
    print(f"server: {len(roll)} student records\n")

    by_reg = {s["registration_number"]: s for s in roll}
    to_change, already, missing = [], [], []
    for sid, name in sorted(wanted.items()):
        student = by_reg.get(sid)
        if not student:
            missing.append((sid, name))
        elif student.get("campus_location") == "LAKESIDE_CAMPUS":
            already.append((sid, name))
        else:
            to_change.append((sid, name, student))

    # Anyone the server has as Lakeside that the sheet does not. Reported, never
    # reverted: the office may have set it deliberately and this script's job is
    # the sheet's Yes list, not the whole truth about campuses.
    unexpected = [s for s in roll
                  if s.get("campus_location") == "LAKESIDE_CAMPUS"
                  and s["registration_number"] not in wanted]

    for sid, name, _ in to_change:
        print(f"  CHANGE   {sid}  {name}")
    for sid, name in already:
        print(f"  already  {sid}  {name}")
    for sid, name in missing:
        print(f"  NO ACCOUNT {sid}  {name}   (not imported - check the import review)")
    for s in unexpected:
        print(f"  ON SERVER ONLY  {s['registration_number']}  {s['name']}   (Lakeside here, not in the sheet)")

    print(f"\nto change {len(to_change)} | already correct {len(already)} | "
          f"no account {len(missing)} | server-only {len(unexpected)}")

    if not apply:
        print("\nDry run. Re-run with --apply to write these changes.")
        return 0
    if not to_change:
        print("\nNothing to do.")
        return 0

    confirm = input(f"\nWrite {len(to_change)} membership changes to PRODUCTION? [type yes] ")
    if confirm.strip().lower() != "yes":
        print("Aborted, nothing written.")
        return 1

    ok = 0
    for sid, name, student in to_change:
        try:
            call("PATCH", f"/admin/students/{student['id']}/membership", token, {
                # student_type is sent unchanged: this script is about campus.
                "student_type": student.get("student_type") or "HOSTELLER",
                "campus_location": "LAKESIDE_CAMPUS",
                "reason": REASON,
            })
            ok += 1
            print(f"  done  {sid}  {name}")
        except SystemExit as e:
            print(f"  FAILED {sid}  {name}: {e}")
    print(f"\n{ok} of {len(to_change)} updated. Every change is in the audit log.")
    return 0 if ok == len(to_change) else 1


if __name__ == "__main__":
    raise SystemExit(main())
