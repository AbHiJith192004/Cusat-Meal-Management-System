"""Opted-in-day billing integration tests in per-test disposable PostgreSQL schemas."""
import asyncio
import io
import os
import uuid
from datetime import datetime, date
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest
import pytest_asyncio
from sqlalchemy import select, text, func
from sqlalchemy.engine import make_url
from sqlalchemy.ext.asyncio import create_async_engine, async_sessionmaker
from sqlalchemy.pool import NullPool

from app.config import get_settings
from app.models import Base
from app.models.user import User
from app.models.meal import MealSelection
from app.models.holiday import Holiday
from app.models.fine import Fine
from app.models.attendance import Attendance
from app.models.billing import StudentBillSnapshot
from app.models.student import StudentProfile
from app.services.billing_service import BillingService
from app.services.student_billing_service import StudentBillingService
from app.services.report_service import ReportService
from app.services.fine_service import FineService
from app.services.holiday_service import HolidayService
from app.services.attendance_service import AttendanceService
from app.services.meal_service import MealService
from app.utils.timezone import IST
from app.utils.exceptions import ConflictException, ValidationException, NotFoundException

YEAR, MONTH = 2024, 5
FIGURES = {'opening_stock_value': '100.00', 'purchases_value': '300.00', 'closing_stock_value': '100.00',
           'operational_expenses': '60.00', 'administrative_expenses': '30.00'}


@pytest_asyncio.fixture
async def database():
    url = make_url(get_settings().DATABASE_URL)
    assert url.host in {'localhost', '127.0.0.1'} and url.database.endswith('_test')
    schema = 'billing_' + uuid.uuid4().hex
    owner = create_async_engine(url, poolclass=NullPool)
    async with owner.begin() as conn:
        await conn.execute(text(f'CREATE SCHEMA {schema}'))
    engine = create_async_engine(url, poolclass=NullPool,
        connect_args={'server_settings': {'search_path': schema}})
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    async with factory() as db:
        enrolled = datetime(2024, 5, 1, tzinfo=IST)
        users = [User(id=uuid.uuid4(), registration_number=name, name=name,
                      role='STUDENT' if name != 'ADMIN' else 'ADMIN', account_status='ACTIVE',
                      created_at=enrolled, activated_at=enrolled) for name in ['A', 'B', 'ADMIN']]
        db.add_all(users)
        await db.commit()
    yield factory, users
    await engine.dispose()
    async with owner.begin() as conn:
        await conn.execute(text(f'DROP SCHEMA {schema} CASCADE'))
    await owner.dispose()


async def preview_publish(factory, admin, figures=None):
    figures = dict(figures or FIGURES)
    async with factory() as db:
        service = BillingService(db)
        preview = await service.preview(MONTH, YEAR, figures)
        result = await service.publish(MONTH, YEAR, {**figures, 'preview_token': preview['preview_token']}, admin.id)
        return preview, result


@pytest.mark.asyncio
async def test_opted_days_partial_skips_holidays_fines_and_zero_attendance(database):
    factory, (a, b, admin) = database
    async with factory() as db:
        # A: full cut May 2, single skip May 3. Both students: no service May 4.
        db.add_all([MealSelection(student_id=a.id, meal_date=date(YEAR, MONTH, 2), meal_type=m, status='SKIPPED')
                    for m in ['BREAKFAST', 'LUNCH', 'DINNER']])
        db.add(MealSelection(student_id=a.id, meal_date=date(YEAR, MONTH, 3), meal_type='BREAKFAST', status='SKIPPED'))
        db.add(Holiday(holiday_date=date(YEAR, MONTH, 4), reason='Closed kitchen', created_by=admin.id))
        db.add(Fine(student_id=a.id, meal_date=date(YEAR, MONTH, 5), meal_type='LUNCH', amount=Decimal('30'), status='PENDING'))
        db.add(Fine(student_id=a.id, meal_date=date(YEAR, MONTH, 6), meal_type='LUNCH', amount=Decimal('50'), status='WAIVED',
                    waived_at=datetime(2024,5,7,tzinfo=IST), waived_by=admin.id, waiver_reason='Verified exemption'))
        await db.commit()
    preview, _ = await preview_publish(factory, admin)
    assert preview['chargeable_days'] == 59
    assert preview['grand_total_expense'] == '390.00'
    bills = preview['students']
    assert bills[0]['effective_days'] == 29 and bills[1]['effective_days'] == 30
    assert '2024-05-03' in bills[0]['opted_in_days']
    assert bills[0]['total_fines'] == '30.00'
    assert sum(Decimal(x['base_charge']) for x in bills) == Decimal('390.00')
    assert sum(Decimal(x['grand_total']) for x in bills) == Decimal('420.00')
    async with factory() as db:
        assert await db.scalar(select(func.count()).select_from(Attendance)) == 0
        assert (await StudentBillingService(db).get_bill(a.id, MONTH, YEAR))['effective_days'] == 29


@pytest.mark.asyncio
async def test_published_snapshot_immutable_and_correction_keeps_old_revision(database):
    factory, (a, b, admin) = database
    await preview_publish(factory, admin)
    async with factory() as db:
        original = await StudentBillingService(db).get_bill(a.id, MONTH, YEAR)
        # Deliberate direct data change simulates an external source correction.
        db.add(Fine(student_id=a.id, meal_date=date(YEAR, MONTH, 5), meal_type='LUNCH', amount=Decimal('30'), status='PENDING'))
        stored = await db.get(User, a.id); stored.name = 'Corrected student name'
        db.add(Attendance(student_id=a.id, meal_date=date(YEAR, MONTH, 5), meal_type='LUNCH', attendance_type='MANUAL', recorded_by=admin.id))
        await db.commit()
    async with factory() as db:
        assert await StudentBillingService(db).get_bill(a.id, MONTH, YEAR) == original
        await BillingService(db).unpublish(MONTH, YEAR, 'Correct missed fine', admin.id)
        with pytest.raises(NotFoundException):
            await StudentBillingService(db).get_bill(a.id, MONTH, YEAR)
    await preview_publish(factory, admin)
    async with factory() as db:
        revised = await StudentBillingService(db).get_bill(a.id, MONTH, YEAR)
        assert revised['revision'] == 2 and revised['total_fines'] == '30.00'
        old = await db.scalar(select(StudentBillSnapshot).where(StudentBillSnapshot.student_id == a.id, StudentBillSnapshot.revision == 1))
        assert old.payload == original


@pytest.mark.asyncio
async def test_stale_preview_cannot_publish(database):
    factory, (a, b, admin) = database
    async with factory() as db:
        preview = await BillingService(db).preview(MONTH, YEAR, FIGURES)
        await db.commit()
    async with factory() as db:
        db.add(MealSelection(student_id=a.id, meal_date=date(YEAR, MONTH, 2), meal_type='LUNCH', status='SKIPPED'))
        db.add(Fine(student_id=a.id, meal_date=date(YEAR, MONTH, 3), meal_type='LUNCH', amount=Decimal('30'), status='PENDING'))
        await db.commit()
    async with factory() as db:
        with pytest.raises(ConflictException):
            await BillingService(db).publish(MONTH, YEAR, {**FIGURES, 'preview_token': preview['preview_token']}, admin.id)
        assert await db.scalar(select(func.count()).select_from(StudentBillSnapshot)) == 0


@pytest.mark.asyncio
async def test_concurrent_publish_and_cent_allocation(database):
    factory, (_, _, admin) = database
    figures = {**FIGURES, 'opening_stock_value': '0', 'purchases_value': '0.01', 'closing_stock_value': '0',
               'operational_expenses': '0', 'administrative_expenses': '0'}
    async with factory() as db:
        preview = await BillingService(db).preview(MONTH, YEAR, figures)
    assert sorted(x['base_charge'] for x in preview['students']) == ['0.00', '0.01']
    async def publish():
        async with factory() as db:
            try:
                await BillingService(db).publish(MONTH, YEAR, {**figures, 'preview_token': preview['preview_token']}, admin.id)
                return 'published'
            except ConflictException:
                return 'conflict'
    assert sorted(await asyncio.gather(publish(), publish())) == ['conflict', 'published']


@pytest.mark.asyncio
async def test_fines_use_implicit_optin_and_block_premature_reconciliation(database):
    factory, (a, b, admin) = database
    async with factory() as db:
        service = FineService(db)
        with pytest.raises(ValidationException):
            await service.reconcile_missed_meals(date(2099, 1, 1), 'LUNCH', admin.id)
        assert await service.reconcile_missed_meals(date(YEAR, MONTH, 2), 'LUNCH', admin.id) == 2
        await db.commit()
    async with factory() as db:
        assert await FineService(db).reconcile_missed_meals(date(YEAR, MONTH, 2), 'LUNCH', admin.id) == 0
        await db.commit()
    await preview_publish(factory, admin)
    async with factory() as db:
        fine = await db.scalar(select(Fine).where(Fine.student_id == a.id))
        with pytest.raises(ConflictException):
            await FineService(db).waive_fine(fine.id, 'Verified correction', admin.id)
        with pytest.raises(ConflictException):
            await AttendanceService(db).record_manual_attendance(
                a.id, date(YEAR, MONTH, 2), 'BREAKFAST', 'MANUAL',
                'Historical correction after publication', admin.id)
        with pytest.raises(ConflictException):
            await MealService(db).update_meal_selection(
                a.id, date(YEAR, MONTH, 2), 'BREAKFAST', 'SKIPPED', a.id)


@pytest.mark.asyncio
async def test_holiday_cancellation_preserves_prior_skip(database):
    factory, (a, b, admin) = database
    async with factory() as db:
        db.add(MealSelection(student_id=a.id, meal_date=date(YEAR, MONTH, 2), meal_type='BREAKFAST', status='SKIPPED'))
        await db.commit()
        holiday = await HolidayService(db).declare_holiday(date(YEAR, MONTH, 2), None, 'Kitchen closure', admin.id)
        await db.commit()
        await HolidayService(db).delete_holiday(holiday.id, admin.id)
        await db.commit()
    async with factory() as db:
        assert await db.scalar(select(MealSelection.status).where(MealSelection.student_id == a.id)) == 'SKIPPED'


@pytest.mark.asyncio
async def test_exports_match_snapshot_and_escape_spreadsheet_formulas(database):
    factory, (a, b, admin) = database
    async with factory() as db:
        obj = await db.get(User, a.id); obj.name = '=1+1'
        await db.commit()
    await preview_publish(factory, admin)
    async with factory() as db:
        bills = await StudentBillingService(db).list_published(MONTH, YEAR)
        report = ReportService(db)
        excel = await report.generate_monthly_excel_report(YEAR, MONTH)
        pdf = await report.generate_monthly_pdf_report(YEAR, MONTH)
    ws = openpyxl.load_workbook(io.BytesIO(excel)).active
    assert ws['B2'].value == "'=1+1"
    assert Decimal(str(ws['G2'].value)) == Decimal(bills[0]['grand_total'])
    assert Decimal(str(ws['G4'].value)) == Decimal('390')
    assert pdf.startswith(b'%PDF')
    # Synthetic exports can be refreshed for visual inspection, but only when
    # asked for. These land in docs/verification/, which is tracked, so writing
    # them unconditionally left every test run - local or CI - with a dirty
    # working tree and two binary files changed by timestamp noise alone.
    if os.environ.get('WRITE_VERIFICATION_ARTIFACTS') == '1':
        output = Path(__file__).resolve().parents[2] / 'docs/verification'
        if output.exists():
            (output / 'billing-synthetic.pdf').write_bytes(pdf)
            (output / 'billing-synthetic.xlsx').write_bytes(excel)


@pytest.mark.asyncio
async def test_midmonth_enrollment_pending_and_zero_denominator(database):
    factory, (a, b, admin) = database
    async with factory() as db:
        ua, ub = await db.get(User, a.id), await db.get(User, b.id)
        ua.activated_at = datetime(YEAR, MONTH, 20, tzinfo=IST)
        ub.account_status = 'PENDING'; ub.activated_at = None
        await db.commit()
        preview = await BillingService(db).preview(MONTH, YEAR, FIGURES)
        assert preview['chargeable_days'] == 12 and len(preview['students']) == 1
        ua.account_status = 'PENDING'; ua.activated_at = None
        await db.commit()
        empty = await BillingService(db).preview(MONTH, YEAR, FIGURES)
        with pytest.raises(ValidationException):
            await BillingService(db).publish(MONTH, YEAR, {**FIGURES, 'preview_token': empty['preview_token']}, admin.id)


@pytest.mark.asyncio
async def test_outmess_students_are_billed_nothing_and_do_not_dilute_the_rate(database):
    """An outmess student takes no meals, so they must not appear on a bill.

    The trap this guards: a missing MealSelection row counts as CONFIRMED, so
    an outmess student included in the cohort would silently draw all 31 days
    of May, be charged a third of the expenses for food they never ate, and
    push the denominator from 62 to 93 -- quietly reducing what A and B owe.
    The exclusion is by student_type, which nothing in billing read before.
    """
    factory, (a, b, admin) = database
    enrolled = datetime(2024, 5, 1, tzinfo=IST)
    async with factory() as db:
        outmess = User(id=uuid.uuid4(), registration_number='OUTMESS1', name='Not On The Mess',
                       role='STUDENT', account_status='ACTIVE',
                       created_at=enrolled, activated_at=enrolled)
        db.add(outmess)
        await db.flush()
        db.add(StudentProfile(id=uuid.uuid4(), user_id=outmess.id, date_of_birth=date(2004, 10, 7),
                              student_type='OUTMESS', campus_location='MAIN_CAMPUS'))
        # A hosteller with a profile, to prove the exclusion keys on the type
        # rather than merely on having a profile row at all.
        db.add(StudentProfile(id=uuid.uuid4(), user_id=a.id, date_of_birth=date(2004, 1, 1),
                              student_type='HOSTELLER', campus_location='MAIN_CAMPUS'))
        await db.commit()

    # Fines first: publishing freezes the period against further changes.
    # A meal they never booked is not a meal they missed, so the count is the
    # two real students and no Fine row names the outmess account.
    async with factory() as db:
        assert await FineService(db).reconcile_missed_meals(date(YEAR, MONTH, 2), 'LUNCH', admin.id) == 2
        await db.commit()
    async with factory() as db:
        assert await db.scalar(select(func.count()).select_from(Fine)
                               .where(Fine.student_id == outmess.id)) == 0

    preview, _ = await preview_publish(factory, admin)

    numbers = [s['student']['registration_number'] for s in preview['students']]
    assert numbers == ['A', 'B'], f'outmess student reached the bill: {numbers}'
    # 31 days x 2 students, not x 3.
    assert preview['chargeable_days'] == 62
    # Every rupee still allocated, just across the students who actually ate.
    # Had the outmess student been included this would be 130.00 each.
    assert preview['students'][0]['base_charge'] == '195.00'
    assert sum(Decimal(x['base_charge']) for x in preview['students']) == Decimal('390.00')

    async with factory() as db:
        assert await db.scalar(
            select(func.count()).select_from(StudentBillSnapshot)
            .where(StudentBillSnapshot.student_id == outmess.id)) == 0
