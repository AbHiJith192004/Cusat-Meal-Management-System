import uuid
from datetime import date, datetime, timedelta
from decimal import Decimal

import pytest
from sqlalchemy import delete, func, select, true as sa_true

from app.database import async_session_factory
from app.models.attendance import Attendance
from app.models.audit import AuditLog
from app.models.billing import BillingPeriod, StudentBillSnapshot
from app.models.holiday import Holiday
from app.models.meal import MealSelection
from app.models.student import StudentProfile
from app.models.operations import InventoryItem, LedgerEntry, MenuPublication, PaymentSubmission
from app.models.notification import Notification
from app.models.user import User
from app.utils.timezone import now_ist
# `client` is a pytest fixture, resolved by name from this module's
# namespace, so importing it is what makes it available - it is not an
# unused import. ruff's F401 removed it once and broke every test here.
from prod_tests.test_hardening import client, headers, user  # noqa: F401


@pytest.mark.asyncio
async def test_operational_workflows_persist_validate_and_audit(client):
    admin, student, second = await user('ADMIN'), await user(), await user()
    h = headers(admin)
    day = now_ist().date()

    menu = await client.put(f'/api/v1/admin/menus/{day}/LUNCH',
        json={'items': ['Rice', 'Dal'], 'notes': 'Vegetarian'}, headers=h)
    assert menu.status_code == 200, menu.text
    listed = await client.get(f'/api/v1/menus?start_date={day}&end_date={day}', headers=headers(student))
    # Pick the row this test wrote. data[0] used to be assumed to be it, but the
    # listing orders by (date, meal_type) and says nothing about how many meals
    # another test may have published for today -- DINNER sorts before LUNCH.
    assert listed.status_code == 200
    lunch = next(m for m in listed.json()['data'] if m['meal_type'] == 'LUNCH')
    assert lunch['items'] == ['Rice', 'Dal']

    ledger = await client.post('/api/v1/admin/ledger', json={'entry_date': str(day),
        'kind': 'PURCHASE', 'category': 'Rice', 'description': 'Monthly rice delivery',
        'amount': '1250.50', 'vendor': 'Synthetic supplier'}, headers=h)
    assert ledger.status_code == 200, ledger.text
    ledger_id = ledger.json()['data']['id']
    assert (await client.post(f'/api/v1/admin/ledger/{ledger_id}/void',
        json={'reason': 'Duplicate supplier invoice'}, headers=h)).status_code == 200

    stock = await client.post('/api/v1/admin/inventory', json={'sku': 'SYN-' + uuid.uuid4().hex[:8],
        'name': 'Synthetic rice', 'unit': 'kg', 'opening_quantity': '10',
        'reorder_level': '3', 'unit_cost': '50'}, headers=h)
    assert stock.status_code == 200, stock.text
    stock_id = stock.json()['data']['id']
    adjusted = await client.post(f'/api/v1/admin/inventory/{stock_id}/adjust',
        json={'quantity_delta': '-8', 'reason': 'Kitchen issue'}, headers=h)
    assert adjusted.status_code == 200 and adjusted.json()['data']['low_stock'] is True
    rejected = await client.post(f'/api/v1/admin/inventory/{stock_id}/adjust',
        json={'quantity_delta': '-3', 'reason': 'Would be negative'}, headers=h)
    assert rejected.status_code == 409

    start, end = now_ist() - timedelta(minutes=1), now_ist() + timedelta(hours=3)
    assigned = await client.post('/api/v1/admin/committee/promote', json={'student_id': str(student.id),
        'starts_at': start.isoformat(), 'ends_at': end.isoformat(), 'scope': 'ATTENDANCE_SCANNER'}, headers=h)
    assert assigned.status_code == 200, assigned.text
    profile = await client.get('/api/v1/me', headers=headers(student))
    assert profile.json()['data']['capabilities']['attendance_scanner'] is True
    assert (await client.post('/api/v1/attendance/verify', json={'qr_token':'invalid'}, headers=headers(second))).status_code == 403
    assert (await client.post('/api/v1/attendance/verify', json={'qr_token':'invalid'}, headers=headers(student))).status_code != 403
    overlap = await client.post('/api/v1/admin/committee/promote', json={'student_id': str(student.id),
        'starts_at': start.isoformat(), 'ends_at': end.isoformat(), 'scope': 'ATTENDANCE_SCANNER'}, headers=h)
    assert overlap.status_code == 409

    bulk = await client.post('/api/v1/admin/attendance/bulk-mark', json={'student_ids': [str(student.id), str(second.id)],
        'meal_date': str(day), 'meal_type': 'DINNER', 'reason': 'Signed paper attendance register'}, headers=h)
    assert bulk.status_code == 200 and bulk.json()['data']['recorded'] == 2
    recent = await client.get('/api/v1/attendance/recent', headers=headers(student))
    assert recent.status_code == 200 and any(x['student_name'] == student.name for x in recent.json()['data'])
    duplicate = await client.post('/api/v1/admin/attendance/bulk-mark', json={'student_ids': [str(student.id)],
        'meal_date': str(day), 'meal_type': 'DINNER', 'reason': 'Signed paper attendance register'}, headers=h)
    assert duplicate.status_code == 409

    async with async_session_factory() as db:
        assert await db.scalar(select(func.count()).select_from(MenuPublication).where(MenuPublication.menu_date == day)) >= 1
        assert await db.scalar(select(func.count()).select_from(LedgerEntry).where(LedgerEntry.id == uuid.UUID(ledger_id), LedgerEntry.voided_at.is_not(None))) == 1
        assert await db.scalar(select(InventoryItem.quantity).where(InventoryItem.id == uuid.UUID(stock_id))) == Decimal('2.000')
        assert await db.scalar(select(func.count()).select_from(Attendance).where(Attendance.meal_date == day, Attendance.meal_type == 'DINNER', Attendance.student_id.in_([student.id, second.id]))) == 2
        assert await db.scalar(select(func.count()).select_from(AuditLog).where(AuditLog.actor_id == admin.id)) >= 6


@pytest.mark.asyncio
async def test_payment_is_bound_to_published_invoice_and_reviewed_once(client):
    admin, student = await user('ADMIN'), await user()
    async with async_session_factory() as db:
        existing = set((await db.execute(select(BillingPeriod.month, BillingPeriod.year))).all())
        month, year = next((m, y) for y in range(2024, 2101) for m in range(1, 13) if (m, y) not in existing)
        period = BillingPeriod(month=month, year=year, is_published=True, revision=1,
            published_at=now_ist(), published_by=admin.id)
        db.add(period); await db.flush()
        db.add(StudentBillSnapshot(period_id=period.id, student_id=student.id, revision=1,
            payload={'grand_total': '432.10', 'student': {'name': student.name,
            'registration_number': student.registration_number}}))
        await db.commit()
    wrong = await client.post('/api/v1/me/payments', json={'month':month,'year':year,
        'utr':'SYNTHETIC123','amount':'400.00'}, headers=headers(student))
    assert wrong.status_code == 422
    submitted = await client.post('/api/v1/me/payments', json={'month':month,'year':year,
        'utr':'SYNTHETIC'+uuid.uuid4().hex[:12].upper(),'amount':'432.10'}, headers=headers(student))
    assert submitted.status_code == 200, submitted.text
    payment_id = submitted.json()['data']['id']
    reviewed = await client.post(f'/api/v1/admin/payments/{payment_id}/review',
        json={'decision':'VERIFIED','note':'Matched against bank statement'}, headers=headers(admin))
    assert reviewed.status_code == 200 and reviewed.json()['data']['status'] == 'VERIFIED'
    blocked_reopen = await client.post('/api/v1/admin/bills/unpublish',
        json={'month': month, 'year': year, 'reason':'Correction after payment'}, headers=headers(admin))
    assert blocked_reopen.status_code == 409
    assert (await client.post(f'/api/v1/admin/payments/{payment_id}/review',
        json={'decision':'REJECTED','note':'Second review prohibited'}, headers=headers(admin))).status_code == 409
    async with async_session_factory() as db:
        row = await db.scalar(select(PaymentSubmission).where(PaymentSubmission.id == uuid.UUID(payment_id)))
        assert row.status == 'VERIFIED' and row.reviewed_by == admin.id
        assert await db.scalar(select(func.count()).select_from(Notification).where(
            Notification.user_id == student.id, Notification.title == 'Payment verified')) == 1


@pytest.mark.asyncio
async def test_two_meal_skip_rejected_but_one_and_three_allowed(client):
    """The real rule is 0, 1 or 3 meals skipped - never exactly 2.

    A deleted unit test used to assert 'max 1 meal opt-out per day', which
    contradicts the implementation: skipping all three is a full-day mess cut
    and is explicitly allowed. It tested a locally defined stub, so it never
    noticed.
    """
    student = await user()
    day = (now_ist().date() + timedelta(days=6)).isoformat()

    first = await client.put(f'/api/v1/meals/{day}/BREAKFAST', json={'status': 'SKIPPED'},
                             headers=headers(student))
    assert first.status_code == 200, first.text

    second = await client.put(f'/api/v1/meals/{day}/LUNCH', json={'status': 'SKIPPED'},
                              headers=headers(student))
    assert second.status_code == 422, second.text

    # Three is a mess cut, and the whole-day route reaches it without ever
    # passing through the invalid two-skip state.
    whole = await client.put(f'/api/v1/meals/{day}', json={'status': 'SKIPPED'},
                             headers=headers(student))
    assert whole.status_code == 200, whole.text

    async with async_session_factory() as db:
        skipped = await db.scalar(select(func.count()).select_from(MealSelection).where(
            MealSelection.student_id == student.id,
            MealSelection.meal_date == date.fromisoformat(day),
            MealSelection.status == 'SKIPPED'))
    assert skipped == 3


@pytest.mark.asyncio
async def test_excel_import_creates_pending_students_and_cannot_grant_a_role(client):
    """Drives the real importer through the real endpoint.

    Replaces a unit test that built a workbook with openpyxl and read it back
    with openpyxl, touching no application code at all.
    """
    import io
    import openpyxl

    super_admin = await user('SUPER_ADMIN')
    reg_ok = 'XL-' + uuid.uuid4().hex[:10].upper()
    reg_dupe = 'XL-' + uuid.uuid4().hex[:10].upper()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(['registration_number', 'name', 'date_of_birth', 'student_type', 'mess_id'])
    ws.append([reg_ok, 'Imported One', '2005-05-15', 'HOSTELLER', 'M-1'])
    ws.append([reg_dupe, 'Imported Two', '2004-11-20', 'DAY_SCHOLAR', 'M-2'])
    # A repeated student id is a RESUBMISSION, not an error: students fill the
    # form again to correct themselves, so the later row wins.
    ws.append([reg_dupe, 'Corrected Two', '2004-11-20', 'DAY_SCHOLAR', 'M-3'])
    ws.append(['', 'Missing Registration', '2004-01-01', 'HOSTELLER', 'M-4'])
    buf = io.BytesIO()
    wb.save(buf)

    response = await client.post(
        '/api/v1/super-admin/students/import',
        files={'file': ('students.xlsx', buf.getvalue(),
                        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')},
        headers=headers(super_admin))
    assert response.status_code == 200, response.text
    summary = response.json()['data']

    assert summary['imported_count'] == 2
    # Only the row with no student id is a rejection; the repeat is a
    # correction and is folded into one account.
    assert summary['skipped_count'] == 1
    # Every rejection names its row so staff can fix the sheet.
    assert [err['row'] for err in summary['errors']] == [5]

    async with async_session_factory() as db:
        imported = await db.scalar(select(User).where(User.registration_number == reg_ok))
        assert imported.role == 'STUDENT'
        assert imported.account_status == 'PENDING'
        assert imported.password_hash is None
        # The later answer is the one kept, and only one account exists.
        corrected = (await db.execute(select(User).where(
            User.registration_number == reg_dupe))).scalars().all()
        assert len(corrected) == 1
        assert corrected[0].name == 'Corrected Two'


@pytest.mark.asyncio
async def test_import_cannot_be_tricked_into_creating_an_admin(client):
    """A role column in the sheet must not escalate anyone.

    The importer hardcodes role=STUDENT at construction, so this is structural
    rather than a validation rule. Worth pinning: it is the single guarantee
    standing between a spreadsheet and a privileged account.
    """
    import io
    import openpyxl

    super_admin = await user('SUPER_ADMIN')
    reg = 'XLROLE-' + uuid.uuid4().hex[:8].upper()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(['registration_number', 'name', 'date_of_birth', 'student_type', 'mess_id', 'role'])
    ws.append([reg, 'Would Be Admin', '2004-02-02', 'HOSTELLER', 'M-9', 'SUPER_ADMIN'])
    buf = io.BytesIO()
    wb.save(buf)

    response = await client.post(
        '/api/v1/super-admin/students/import',
        files={'file': ('students.xlsx', buf.getvalue(),
                        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')},
        headers=headers(super_admin))
    assert response.status_code == 200, response.text

    async with async_session_factory() as db:
        created = await db.scalar(select(User).where(User.registration_number == reg))
    assert created.role == 'STUDENT'


@pytest.mark.asyncio
async def test_monthly_mess_cut_limit_is_driven_by_the_setting(client):
    """Changing max_monthly_mess_cuts must actually change what students hit.

    The limit was a hardcoded 10 in meal_service while this setting existed and
    was read nowhere, so editing it on the settings screen did nothing. Setting
    it to 1 here proves the configured value is the enforced one; a hardcoded
    10 would let the second mess cut through.
    """
    from app.models.settings import SystemSetting

    student, super_admin = await user(), await user('SUPER_ADMIN')
    # Two days in one calendar month, comfortably past the selection cutoff.
    first_of_next = (now_ist().date().replace(day=1) + timedelta(days=32)).replace(day=1)
    day_one, day_two = first_of_next.replace(day=10), first_of_next.replace(day=11)

    async def set_limit(value: str):
        return await client.put('/api/v1/admin/settings',
                                json={'settings': [{'key': 'max_monthly_mess_cuts', 'value': value}]},
                                headers=headers(super_admin))

    assert (await set_limit('1')).status_code == 200
    try:
        first = await client.put(f'/api/v1/meals/{day_one.isoformat()}',
                                 json={'status': 'SKIPPED'}, headers=headers(student))
        assert first.status_code == 200, first.text

        second = await client.put(f'/api/v1/meals/{day_two.isoformat()}',
                                  json={'status': 'SKIPPED'}, headers=headers(student))
        assert second.status_code == 422, second.text
        # The refusal quotes the configured limit, not a baked-in 10.
        assert '1 per month' in second.text

        async with async_session_factory() as db:
            cut_days = await db.scalar(select(func.count(func.distinct(MealSelection.meal_date))).where(
                MealSelection.student_id == student.id, MealSelection.status == 'SKIPPED'))
        assert cut_days == 1
    finally:
        # Settings are global to the database this suite shares, so a low limit
        # left behind would fail unrelated tests.
        async with async_session_factory() as db:
            row = await db.scalar(select(SystemSetting).where(
                SystemSetting.key == 'max_monthly_mess_cuts'))
            if row:
                row.value = '10'
                await db.commit()


def _workbook(rows: list[tuple]) -> bytes:
    """A sheet shaped like the live Google Form export, Timestamp first."""
    import io
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(['Timestamp', 'Student id', 'Full name ', 'Date of birth ', 'Email address ',
               'Phone number ', 'Department ', 'Lakeside ', 'Guest / inmate',
               'Course/programme ', 'Room number ', 'Profile picture '])
    for r in rows:
        # openpyxl cannot write a tz-aware datetime, and a Forms export never
        # contains one - the Timestamp column is naive local time.
        ws.append([v.replace(tzinfo=None) if isinstance(v, datetime) else v for v in r])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.mark.asyncio
async def test_import_never_disturbs_an_account_that_already_exists(client):
    """Re-importing the sheet must not reset a password or revive an account.

    Staff re-upload a corrected workbook routinely, and by then some students
    have already activated. Overwriting them would lock those students out.
    """
    from app.security.password import hash_password

    super_admin = await user('SUPER_ADMIN')
    reg = 'IMP-' + uuid.uuid4().hex[:8].upper()
    sheet = _workbook([
        (now_ist(), reg, 'Imported Student', date(2004, 5, 1), f'{reg.lower()}@example.com',
         9876543210, 'DCA', 'Yes', 'Inmate', 'MCA', '29B', ''),
    ])
    files = {'file': ('students.xlsx', sheet,
                      'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}

    first = await client.post('/api/v1/super-admin/students/import', files=files,
                              headers=headers(super_admin))
    assert first.status_code == 200, first.text
    assert first.json()['data']['imported_count'] == 1

    # The student activates, then staff re-upload the same workbook.
    async with async_session_factory() as db:
        row = await db.scalar(select(User).where(User.registration_number == reg))
        row.password_hash = hash_password('student-chose-this')
        row.account_status = 'ACTIVE'
        await db.commit()

    second = await client.post('/api/v1/super-admin/students/import', files=files,
                               headers=headers(super_admin))
    assert second.status_code == 200, second.text
    body = second.json()['data']
    assert body['imported_count'] == 0
    assert 'already exists' in body['errors'][0]['error'].lower()

    async with async_session_factory() as db:
        after = await db.scalar(select(User).where(User.registration_number == reg))
        duplicates = await db.scalar(select(func.count()).select_from(User).where(
            User.registration_number == reg))
    assert after.account_status == 'ACTIVE'
    assert after.password_hash is not None
    assert duplicates == 1


@pytest.mark.asyncio
async def test_import_maps_the_form_columns_and_types_outmess_students(client):
    """Columns are matched by heading, so the leading Timestamp is harmless.

    The outmess student is imported like anyone else and distinguished by
    student_type. They used to be skipped, which left them with no account
    to join the mess from later; billing and fines exclude them instead,
    which prod_tests/test_billing.py proves.
    """
    super_admin = await user('SUPER_ADMIN')
    keep = 'IMPK-' + uuid.uuid4().hex[:8].upper()
    out = 'IMPD-' + uuid.uuid4().hex[:8].upper()
    sheet = _workbook([
        (now_ist(), keep, ' Padded Name ', date(2003, 2, 14), f'{keep.lower()}@example.com',
         '+919876543211', 'DCS', 'Yes', 'Inmate', 'M.Tech', '58 A', ''),
        (now_ist(), out, 'Outmess Person', date(2003, 3, 3), f'{out.lower()}@example.com',
         9876543212, 'DCA', 'No', 'Outmess', 'MCA', 'NA', ''),
    ])
    response = await client.post(
        '/api/v1/super-admin/students/import',
        files={'file': ('students.xlsx', sheet,
                        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')},
        headers=headers(super_admin))
    assert response.status_code == 200, response.text
    body = response.json()['data']
    assert body['imported_count'] == 2
    assert body['errors'] == []

    async with async_session_factory() as db:
        created = await db.scalar(select(User).where(User.registration_number == keep))
        profile = await db.scalar(select(StudentProfile).where(StudentProfile.user_id == created.id))
        outmess = await db.scalar(select(User).where(User.registration_number == out))
        assert outmess is not None
        outmess_profile = await db.scalar(
            select(StudentProfile).where(StudentProfile.user_id == outmess.id))
        assert outmess_profile.student_type == 'OUTMESS'
        assert profile.student_type == 'HOSTELLER'
    assert created.name == 'Padded Name'
    assert created.phone == '9876543211'          # +91 stripped
    assert created.role == 'STUDENT' and created.account_status == 'PENDING'
    assert created.password_hash is None
    assert profile.room_number == '58 A'          # block letter preserved
    assert profile.hostel_name == 'Lakeside'
    assert profile.date_of_birth == date(2003, 2, 14)


@pytest.mark.asyncio
async def test_import_rejects_a_sheet_whose_headings_do_not_match(client):
    """A wrong file should be refused outright, not imported as junk."""
    import io
    import openpyxl

    super_admin = await user('SUPER_ADMIN')
    wb = openpyxl.Workbook()
    wb.active.append(['Some', 'Unrelated', 'Spreadsheet'])
    wb.active.append([1, 2, 3])
    buf = io.BytesIO()
    wb.save(buf)

    response = await client.post(
        '/api/v1/super-admin/students/import',
        files={'file': ('wrong.xlsx', buf.getvalue(),
                        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')},
        headers=headers(super_admin))
    assert response.status_code == 422
    assert 'missing required columns' in response.text.lower()


@pytest.mark.asyncio
async def test_a_menu_cannot_be_written_for_a_day_that_has_passed(client):
    """Today and tomorrow are plans; yesterday is a record of what was served.

    Rewriting a past day makes the published history disagree with the meal
    students actually ate, and nothing is gained by allowing it.
    """
    admin = await user('ADMIN')
    h = headers(admin)
    today = now_ist().date()

    ok = await client.put(f'/api/v1/admin/menus/{today}/DINNER',
                          json={'items': ['Chapati', 'Curry'], 'notes': None}, headers=h)
    assert ok.status_code == 200, ok.text

    ahead = await client.put(f'/api/v1/admin/menus/{today + timedelta(days=3)}/DINNER',
                             json={'items': ['Rice', 'Sambar'], 'notes': None}, headers=h)
    assert ahead.status_code == 200, ahead.text

    past = await client.put(f'/api/v1/admin/menus/{today - timedelta(days=1)}/DINNER',
                            json={'items': ['Rewritten'], 'notes': None}, headers=h)
    assert past.status_code == 422
    assert past.json()['error']['code'] == 'MENU_DATE_IN_PAST'

    # ... and the refusal wrote nothing. Asserting a COUNT of zero for
    # yesterday was wrong: this same test writes today/DINNER every run, so the
    # day after any earlier run there is legitimately a row there and the
    # assertion failed on its own history rather than on a defect. What
    # actually matters is that the rejected content never landed.
    async with async_session_factory() as db:
        row = await db.scalar(select(MenuPublication).where(
            MenuPublication.menu_date == today - timedelta(days=1),
            MenuPublication.meal_type == 'DINNER'))
        assert row is None or row.items != ['Rewritten']


@pytest.mark.asyncio
async def test_membership_change_is_audited_and_rejects_a_no_op(client):
    """Category and campus both move money, so the reason is kept with them."""
    admin, student = await user('ADMIN'), await user()
    async with async_session_factory() as db:
        db.add(StudentProfile(user_id=student.id, mess_id=f'M-{uuid.uuid4().hex[:8]}',
                              date_of_birth=date(2004, 5, 17), student_type='HOSTELLER',
                              campus_location='MAIN_CAMPUS'))
        await db.commit()

    url = f'/api/v1/admin/students/{student.id}/membership'
    body = {'student_type': 'OUTMESS', 'campus_location': 'LAKESIDE_CAMPUS',
            'reason': 'Moved off the mess from this month, confirmed with the warden'}

    changed = await client.patch(url, json=body, headers=headers(admin))
    assert changed.status_code == 200, changed.text
    assert changed.json()['data']['student_type'] == 'OUTMESS'

    async with async_session_factory() as db:
        profile = await db.scalar(select(StudentProfile).where(StudentProfile.user_id == student.id))
        assert profile.student_type == 'OUTMESS'
        assert profile.campus_location == 'LAKESIDE_CAMPUS'
        entry = await db.scalar(select(AuditLog).where(AuditLog.action == 'STUDENT_MEMBERSHIP_CHANGED',
                                                       AuditLog.actor_id == admin.id))
        assert entry is not None
        assert entry.metadata_['before']['student_type'] == 'HOSTELLER'
        assert entry.metadata_['after']['campus_location'] == 'LAKESIDE_CAMPUS'
        assert 'warden' in entry.metadata_['reason']

    # Re-sending the same values is a mistake, not a change.
    again = await client.patch(url, json=body, headers=headers(admin))
    assert again.status_code == 422

    # A reason is not optional.
    assert (await client.patch(url, headers=headers(admin), json={
        'student_type': 'HOSTELLER', 'campus_location': 'MAIN_CAMPUS', 'reason': 'x',
    })).status_code == 422

    # Students cannot reclassify themselves.
    assert (await client.patch(url, json=body, headers=headers(student))).status_code == 403


@pytest.mark.asyncio
async def test_meal_window_settings_reject_anything_the_app_cannot_parse(client):
    """The settings endpoint is now reachable from a screen, so it has to hold.

    MealTimingService parses these with `map(int, value.split(':'))` inside
    GET /api/v1/meals -- the call every student's home screen makes. A stored
    "12.00" would raise there, so meal selection for the whole hostel depends
    on nothing invalid getting written here.

    System settings are global, so this restores every row it touched. An
    earlier draft did not, left dinner ending at 06:00 behind, and the next
    run failed on its own leftovers before reaching its first assertion.
    """
    from app.models.settings import SystemSetting

    super_admin = await user('SUPER_ADMIN')
    h = headers(super_admin)

    def put(*pairs):
        return client.put('/api/v1/admin/settings',
                          json={'settings': [{'key': k, 'value': v} for k, v in pairs]},
                          headers=h)

    async with async_session_factory() as db:
        before = {s.key: s.value for s in (await db.execute(select(SystemSetting))).scalars()}

    try:
        # A good change is applied and readable back.
        ok = await put(('meal_window_lunch_start', '12:15'), ('meal_window_lunch_end', '14:45'))
        assert ok.status_code == 200, ok.text
        listed = await client.get('/api/v1/admin/settings', headers=h)
        values = {s['key']: s['value'] for s in listed.json()['data']}
        assert values['meal_window_lunch_start'] == '12:15'
        assert values['meal_window_lunch_end'] == '14:45'

        # ...and the students' own endpoint still parses, which is the point.
        student = await user()
        day = now_ist().date()
        meals = await client.get(f'/api/v1/meals?start_date={day}&end_date={day}',
                                 headers=headers(student))
        assert meals.status_code == 200, meals.text
        assert meals.json()['data'][0]['lunch']['time_window'] == '12:15\u201314:45 IST'

        # The mess-cut limit rides along on the same payload so the meal
        # planner can print the real number instead of a hardcoded 10.
        assert (await put(('max_monthly_mess_cuts', '6'))).status_code == 200
        again = await client.get(f'/api/v1/meals?start_date={day}&end_date={day}',
                                 headers=headers(student))
        assert again.json()['data'][0]['max_monthly_mess_cuts'] == '6'

        for bad in ('12.00', 'noon', '25:00', '12:60', '', '9'):
            rejected = await put(('meal_window_lunch_start', bad))
            assert rejected.status_code == 422, f'{bad!r} was accepted: {rejected.text}'

        # An end before its start is refused even though both are valid clock
        # times. Dinner has no stored row here, so this also proves the
        # comparison resolves against the app's own DEFAULT_SETTINGS rather
        # than skipping a window whose other end is absent.
        inverted = await put(('meal_window_dinner_end', '06:00'))
        assert inverted.status_code == 422, inverted.text
        assert inverted.json()['error']['code'] == 'SETTING_WINDOW_INVERTED'

        # A batch moving both ends together is fine.
        both = await put(('meal_window_dinner_start', '05:00'), ('meal_window_dinner_end', '06:00'))
        assert both.status_code == 200, both.text

        # Nothing from a rejected batch is written: lunch start survived.
        after = await client.get('/api/v1/admin/settings', headers=h)
        assert {s['key']: s['value'] for s in after.json()['data']}['meal_window_lunch_start'] == '12:15'

        # The other parsed settings are guarded too.
        assert (await put(('max_monthly_mess_cuts', 'ten'))).status_code == 422
        assert (await put(('fine_amount', '-5'))).status_code == 422
        assert (await put(('selection_cutoff_time', '9pm'))).status_code == 422
        assert (await put(('qr_validity_seconds', '2'))).status_code == 422
        # Two minutes is the ceiling: nobody checks a face at the door, so the
        # pass's short life is what stops a forwarded screenshot working.
        assert (await put(('qr_validity_seconds', '300'))).status_code == 422

        # And the value reaches a real pass, which is the whole point -- the row
        # existed for months while QRService read the environment instead.
        # Lunch is opened to the whole day first: a pass is only issued inside
        # the serving window, and this must not pass or fail on the clock the
        # suite happens to run at. Both settings are restored in `finally`.
        assert (await put(('meal_window_lunch_start', '00:00'),
                          ('meal_window_lunch_end', '23:59'),
                          ('qr_validity_seconds', '90'))).status_code == 200
        pass_now = await client.get('/api/v1/attendance/qr?meal_type=LUNCH',
                                    headers=headers(student))
        assert pass_now.status_code == 200, pass_now.text
        assert pass_now.json()['data']['validity_seconds'] == 90

        # A key this guard does not know is still passed through, as before.
        assert (await put(('some_future_setting', 'anything at all'))).status_code == 200

        # Settings are no longer Super Admin only -- see
        # test_settings_are_open_to_any_admin_but_not_to_students for that
        # boundary. What this test guards is that WIDENING WHO MAY WRITE did
        # not widen WHAT MAY BE WRITTEN: an ordinary admin gets the same
        # validation, not a way round it.
        plain = await user('ADMIN')
        refused = await client.put('/api/v1/admin/settings',
                                   json={'settings': [{'key': 'meal_window_lunch_start',
                                                       'value': 'half twelve'}]},
                                   headers=headers(plain))
        assert refused.status_code == 422, refused.text
    finally:
        async with async_session_factory() as db:
            for row in (await db.execute(select(SystemSetting))).scalars():
                if row.key in before:
                    row.value = before[row.key]
                else:
                    await db.delete(row)
            await db.commit()


@pytest.mark.asyncio
async def test_suspending_a_student_locks_them_out_and_reinstating_respects_activation(client):
    """SUSPENDED was enforced in three places and writable from none of them.

    So this checks the enforcement actually engages, not merely that a column
    changed: the student's existing token must stop working the moment they
    are suspended.
    """
    admin, student = await user('ADMIN'), await user()
    h = headers(admin)
    token = headers(student)

    # Works before.
    assert (await client.get('/api/v1/me', headers=token)).status_code == 200

    suspend = await client.patch(f'/api/v1/admin/students/{student.id}/status',
                                 json={'suspend': True, 'reason': 'Left the hostel mid-term'},
                                 headers=h)
    assert suspend.status_code == 200, suspend.text
    assert suspend.json()['data']['account_status'] == 'SUSPENDED'

    # The token they were already holding is dead, not merely their next login.
    denied = await client.get('/api/v1/me', headers=token)
    assert denied.status_code == 403, denied.text
    assert denied.json()['error']['code'] == 'ACCOUNT_SUSPENDED'

    # Suspending twice is a no-op and says so rather than re-writing.
    again = await client.patch(f'/api/v1/admin/students/{student.id}/status',
                               json={'suspend': True, 'reason': 'Same reason again'}, headers=h)
    assert again.status_code == 422

    # A reason is mandatory, as it is for a membership change.
    bare = await client.patch(f'/api/v1/admin/students/{student.id}/status',
                              json={'suspend': False, 'reason': 'x'}, headers=h)
    assert bare.status_code == 422

    # This student never activated, so reinstating returns them to PENDING --
    # ACTIVE would describe an account with no password as ready to use.
    back = await client.patch(f'/api/v1/admin/students/{student.id}/status',
                              json={'suspend': False, 'reason': 'Readmitted for the new term'},
                              headers=h)
    assert back.status_code == 200, back.text
    assert back.json()['data']['account_status'] == 'PENDING'

    # An activated student comes back ACTIVE instead.
    active = await user()
    async with async_session_factory() as db:
        row = await db.get(User, active.id)
        row.activated_at = now_ist()
        await db.commit()
    await client.patch(f'/api/v1/admin/students/{active.id}/status',
                       json={'suspend': True, 'reason': 'Suspended pending review'}, headers=h)
    restored = await client.patch(f'/api/v1/admin/students/{active.id}/status',
                                  json={'suspend': False, 'reason': 'Review closed, no action'},
                                  headers=h)
    assert restored.json()['data']['account_status'] == 'ACTIVE'

    # Reinstating someone who is not suspended is refused.
    assert (await client.patch(f'/api/v1/admin/students/{active.id}/status',
                               json={'suspend': False, 'reason': 'Already back again'},
                               headers=h)).status_code == 422

    # Both directions are audited with the before/after and the reason.
    async with async_session_factory() as db:
        logs = (await db.execute(select(AuditLog).where(
            AuditLog.action.in_(['STUDENT_SUSPENDED', 'STUDENT_REINSTATED'])))).scalars().all()
    mine = [x for x in logs if x.metadata_.get('student_id') == str(student.id)]
    assert {x.action for x in mine} == {'STUDENT_SUSPENDED', 'STUDENT_REINSTATED'}
    assert any(x.metadata_['reason'] == 'Left the hostel mid-term' for x in mine)

    # An admin account cannot be suspended through the student endpoint.
    other_admin = await user('ADMIN')
    assert (await client.patch(f'/api/v1/admin/students/{other_admin.id}/status',
                               json={'suspend': True, 'reason': 'Should not be possible here'},
                               headers=h)).status_code == 404

    # And a student cannot suspend anyone.
    assert (await client.patch(f'/api/v1/admin/students/{active.id}/status',
                               json={'suspend': True, 'reason': 'Student attempting this'},
                               headers=headers(await user()))).status_code == 403


@pytest.mark.asyncio
async def test_closing_the_mess_for_a_day_is_visible_cancellable_and_keeps_attendance_out(client):
    """A mess closure, walked the way an admin actually walks it.

    The machinery for this existed from the beginning and was unreachable.
    Declaring a holiday cascaded every selection to NO_SERVICE, blocked the
    scanner, suppressed the missed-meal fine and dropped the day out of the
    billing denominator -- but nothing listed a closure, so the id the DELETE
    needs could only be recovered by querying the database by hand, and no
    screen called any of it.
    """
    admin, student = await user('ADMIN'), await user()
    h, s = headers(admin), headers(student)
    # A fixed far-future day. Every other test works near today, and a
    # whole-day closure cascades to every selection sharing its date.
    day = date(2031, 3, 17)
    async with async_session_factory() as db:
        await db.execute(delete(Holiday).where(Holiday.holiday_date == day))
        await db.commit()

    # The student opts out of breakfast first, so the revert at the end has
    # something to restore that is not just "CONFIRMED".
    assert (await client.put(f'/api/v1/meals/{day}/BREAKFAST',
                             json={'status': 'SKIPPED'}, headers=s)).status_code == 200

    declared = await client.post('/api/v1/admin/holidays', json={
        'date': str(day), 'meal_type': None, 'reason': 'Kitchen closed for repairs'}, headers=h)
    assert declared.status_code == 200, declared.text

    # The student is told, and told why -- not left with a greyed-out day.
    plan = await client.get(f'/api/v1/meals?start_date={day}&end_date={day}', headers=s)
    assert plan.status_code == 200, plan.text
    today_plan = plan.json()['data'][0]
    for meal in ('breakfast', 'lunch', 'dinner'):
        assert today_plan[meal]['status'] == 'NO_SERVICE'
        assert today_plan[meal]['no_service_reason'] == 'Kitchen closed for repairs'

    # ...and the closure can be found again, which is the part that was missing.
    listed = await client.get(f'/api/v1/admin/holidays?start={day}&end={day}', headers=h)
    assert listed.status_code == 200, listed.text
    rows = listed.json()['data']
    assert len(rows) == 1 and rows[0]['meal_type'] is None
    assert rows[0]['reason'] == 'Kitchen closed for repairs'
    assert rows[0]['declared_by'] == admin.name
    holiday_id = rows[0]['id']

    # Marking attendance by hand used to sail straight past the closure, which
    # left attendance rows for a meal the bill says was never served.
    manual = await client.post('/api/v1/admin/attendance/manual', json={
        'student_id': str(student.id), 'meal_date': str(day), 'meal_type': 'LUNCH',
        'attendance_type': 'MANUAL', 'reason': 'Ate at the counter'}, headers=h)
    assert manual.status_code == 422, manual.text
    assert 'closed' in manual.json()['error']['message'].lower()

    # The pass says why rather than "attendance cannot be recorded at this
    # time", which is also what being early says.
    from app.services.qr_service import QRService
    from app.utils.exceptions import AttendanceUnavailableException
    async with async_session_factory() as db:
        with pytest.raises(AttendanceUnavailableException) as refused:
            await QRService(db)._check_eligibility(student.id, day, 'LUNCH')
    assert 'Kitchen closed for repairs' in refused.value.message

    # The dashboard counted closures as a literal zero before this.
    dashboard = await client.get('/api/v1/admin/dashboard', headers=h)
    assert dashboard.json()['data']['active_holidays_count'] >= 1

    # Cancelling restores what each student had chosen, not a blanket opt-in.
    assert (await client.delete(f'/api/v1/admin/holidays/{holiday_id}', headers=h)).status_code == 200
    reopened = (await client.get(f'/api/v1/meals?start_date={day}&end_date={day}',
                                 headers=s)).json()['data'][0]
    assert reopened['breakfast']['status'] == 'SKIPPED'
    assert reopened['lunch']['status'] == 'CONFIRMED'
    assert reopened['breakfast']['no_service_reason'] is None
    assert (await client.get(f'/api/v1/admin/holidays?start={day}&end={day}',
                             headers=h)).json()['data'] == []

    # A single sitting can be shut without touching the rest of the day.
    one = await client.post('/api/v1/admin/holidays', json={
        'date': str(day), 'meal_type': 'DINNER', 'reason': 'Hostel day feast outside'}, headers=h)
    assert one.status_code == 200, one.text
    partial = (await client.get(f'/api/v1/meals?start_date={day}&end_date={day}',
                                headers=s)).json()['data'][0]
    assert partial['dinner']['status'] == 'NO_SERVICE'
    assert partial['dinner']['no_service_reason'] == 'Hostel day feast outside'
    assert partial['lunch']['status'] == 'CONFIRMED'
    assert partial['breakfast']['status'] == 'SKIPPED'

    # Leave the shared database as it was found.
    assert (await client.delete(
        f"/api/v1/admin/holidays/{one.json()['data']['id']}", headers=h)).status_code == 200
    async with async_session_factory() as db:
        await db.execute(delete(MealSelection).where(MealSelection.meal_date == day))
        await db.commit()


@pytest.mark.asyncio
async def test_settings_are_open_to_any_admin_but_not_to_students(client):
    """Settings moved off the Super Admin router, so prove the new line holds.

    The screen was Super Admin only and so were the endpoints. Both were
    widened together -- a button that 403s is worse than no button. What must
    NOT widen with them is the rest of /super-admin: creating administrators
    and importing students still decide who holds power.
    """
    from app.models.settings import SystemSetting
    admin, student, super_admin = await user('ADMIN'), await user(), await user('SUPER_ADMIN')

    async with async_session_factory() as db:
        before = {s.key: s.value for s in (await db.execute(select(SystemSetting))).scalars()}

    try:
        # A plain admin can read them...
        listed = await client.get('/api/v1/admin/settings', headers=headers(admin))
        assert listed.status_code == 200, listed.text

        # ...and write them, with the same validation everyone else gets.
        written = await client.put('/api/v1/admin/settings', json={'settings': [
            {'key': 'fine_amount', 'value': '35.00'}]}, headers=headers(admin))
        assert written.status_code == 200, written.text
        back = await client.get('/api/v1/admin/settings', headers=headers(admin))
        assert {s['key']: s['value'] for s in back.json()['data']}['fine_amount'] == '35.00'

        # The guard rails did not move with the door.
        bad = await client.put('/api/v1/admin/settings', json={'settings': [
            {'key': 'meal_window_lunch_start', 'value': 'noon'}]}, headers=headers(admin))
        assert bad.status_code == 422, bad.text

        # A student cannot reach either verb.
        assert (await client.get('/api/v1/admin/settings',
                                 headers=headers(student))).status_code == 403
        assert (await client.put('/api/v1/admin/settings', json={'settings': [
            {'key': 'fine_amount', 'value': '0.00'}]},
            headers=headers(student))).status_code == 403

        # The change is attributable: the audit names the admin who made it.
        async with async_session_factory() as db:
            logs = (await db.execute(select(AuditLog).where(
                AuditLog.actor_id == admin.id))).scalars().all()
        assert any('fine_amount' in (x.metadata_ or {}).get('updated_keys', []) for x in logs)

        # Still Super Admin only, and still refused to a plain admin.
        assert (await client.post('/api/v1/super-admin/admins', json={
            'registration_number': 'NOTALLOWED' + uuid.uuid4().hex[:8].upper(),
            'name': 'Should not be created', 'role': 'ADMIN'},
            headers=headers(admin))).status_code == 403
        assert (await client.post('/api/v1/admin/students', json={
            'registration_number': 'NOPE' + uuid.uuid4().hex[:8].upper(),
            'name': 'Should not be created', 'student_type': 'HOSTELLER',
            'campus_location': 'MAIN_CAMPUS'},
            headers=headers(admin))).status_code == 403
    finally:
        # Leave the shared database exactly as it was found.
        restore = [{'key': k, 'value': v} for k, v in before.items()]
        if restore:
            await client.put('/api/v1/admin/settings', json={'settings': restore},
                             headers=headers(super_admin))
        async with async_session_factory() as db:
            await db.execute(delete(SystemSetting).where(
                SystemSetting.key.not_in(list(before)) if before else sa_true()))
            await db.commit()
